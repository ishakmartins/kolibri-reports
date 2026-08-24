"""CMS build step: cms/reports.xlsx -> src/data/reports.json, and stage
reports/<slug>/ into public/<slug>/ so Astro copies them verbatim. <slug> is the
"Report Page link" column, which is what the homepage links to (not the DOI).

Run before `astro build` (npm run build does it).
"""
import json
import shutil
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
COLUMNS = ["UUID", "Title", "Date published", "Report Page link", "DOI", "Kolibri version"]
KEYS = ["uuid", "title", "date", "reportLink", "doi", "version"]
EXCEL_EPOCH = date(1899, 12, 30)  # Excel serial 1 == 1900-01-01, with the 1900 leap bug


def norm_date(v):
    """Excel serial, datetime, or string -> ISO YYYY-MM-DD. Unparseable passes through."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return (EXCEL_EPOCH + timedelta(days=int(v))).isoformat()
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d %B %Y", "%B %d, %Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return s  # ponytail: unknown formats sort lexically; add a parser when one shows up


def read_rows(xlsx):
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, data_only=True)["reports"]
    header = [str(c.value).strip() if c.value else "" for c in ws[1]]
    if header[: len(COLUMNS)] != COLUMNS:
        sys.exit(f"cms: sheet 'reports' header must be {COLUMNS}, got {header}")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        row = dict(zip(KEYS, [("" if v is None else str(v).strip()) for v in r[: len(KEYS)]]))
        row["date"] = norm_date(r[2])
        rows.append(row)
    rows.sort(key=lambda x: x["date"], reverse=True)
    return rows


def stage_reports(src, out):
    slugs = [d for d in sorted(src.glob("*")) if d.is_dir()]
    for d in slugs:
        shutil.rmtree(out / d.name, ignore_errors=True)
        shutil.copytree(d, out / d.name)
    return [d.name for d in slugs]


def main():
    rows = read_rows(ROOT / "cms" / "reports.xlsx")
    (ROOT / "src" / "data").mkdir(parents=True, exist_ok=True)
    (ROOT / "src" / "data" / "reports.json").write_text(json.dumps(rows, indent=2), encoding="utf-8")
    slugs = stage_reports(ROOT / "reports", ROOT / "public")
    print(f"cms: {len(rows)} rows, {len(slugs)} report folders {slugs}")


def test():
    assert norm_date(45000) == "2023-03-15", norm_date(45000)
    assert norm_date(datetime(2026, 2, 1, 9, 30)) == "2026-02-01"
    assert norm_date("2026-02-01") == "2026-02-01"
    assert norm_date("01/02/2026") == "2026-02-01"
    assert norm_date("14 March 2026") == "2026-03-14"
    assert norm_date(None) == ""
    print("ok")


if __name__ == "__main__":
    test() if "--test" in sys.argv else main()
